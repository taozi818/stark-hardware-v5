# -*- coding: utf-8 -*-
import os, re, html, zipfile

XLSX = r"C:\Users\董梅\AccioWork\2026-09-10-19-27-28-237-8204b69a\temp_data_package\资料准备说明\产品类目图片--中英文名字都对应下\产品图.xlsx"
OUT = r"C:\Users\董梅\AccioWork\2026-09-10-19-27-28-237-8204b69a\src\en\products.html"


def read_rows(path):
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        out = []
        for r in ws.iter_rows(values_only=True):
            r = list(r) + ['', '', '', '']
            out.append({'A': r[0], 'B': r[1], 'C': r[2], 'D': r[3]})
        return out
    except Exception:
        pass
    rows = []
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        shared = []
        if 'xl/sharedStrings.xml' in names:
            s = z.read('xl/sharedStrings.xml').decode('utf-8', 'replace')
            for si in re.findall(r'<si>(.*?)</si>', s, re.S):
                txt = ''.join(re.findall(r'<t[^>]*>(.*?)</t>', si, re.S))
                shared.append(html.unescape(txt))
        sheets = sorted(n for n in names if re.match(r'xl/worksheets/sheet\d+\.xml$', n))
        sheet = z.read(sheets[0]).decode('utf-8', 'replace')
        for row in re.findall(r'<row[^>]*>(.*?)</row>', sheet, re.S):
            cells = {}
            for m in re.finditer(r'<c\b([^>]*?)(?:/>|>(.*?)</c>)', row, re.S):
                attrs = m.group(1) or ''
                body = m.group(2) or ''
                ref = re.search(r'r="([A-Z]+)\d+"', attrs)
                col = ref.group(1) if ref else 'A'
                t = re.search(r't="([^"]+)"', attrs)
                typ = t.group(1) if t else ''
                val = ''
                if typ == 'inlineStr':
                    val = ''.join(re.findall(r'<t[^>]*>(.*?)</t>', body, re.S))
                else:
                    v = re.search(r'<v>(.*?)</v>', body, re.S)
                    if v:
                        val = v.group(1)
                        if typ == 's':
                            try:
                                val = shared[int(val)]
                            except Exception:
                                pass
                cells[col] = html.unescape(val)
            rows.append(cells)
    return rows


def esc(s):
    s = str(s or '').replace('&amp;', '&').strip()
    return html.escape(s, quote=True)


GROUPS = [
    ("delay valve", "Commercial Delay Valve Series",
     "Heavy-duty brass self-closing and time-delay flush valves for public restrooms, hospitals and schools",
     "https://s.alicdn.com/@sc04/kf/Hd87b785020ee4b56baa9b24c1d2957c4g/Double-Outlet-Garden-Tap-1-2-3.jpg"),
    ("Induction Sanitary Ware Series", "Induction Sanitary Ware Series",
     "Touchless infrared sensor faucets, urinal flushers and automatic soap dispensers",
     "https://s.alicdn.com/@sc04/kf/H9cad3bf47aaf46c79ad271e76469a458j/Chrome-Automatic-Sensor-Bathroom-Basin-Faucet-Touchless.jpg"),
    ("304 stainless steel commercial series", "304 Stainless Steel Commercial Series",
     "SUS304 vandal-resistant sanitary ware for public, hotel, school and hospital projects",
     "https://s.alicdn.com/@sc04/kf/H99fbb47d32ed491d8a3267928eac2de3h/SUS304-Stainless-Steel-One-Piece-Floor-Mounted.jpg"),
]

rows = read_rows(XLSX)
items = []
for r in rows[1:]:
    cat = str(r.get('B') or '').strip()
    title = str(r.get('C') or '').strip()
    if cat and title:
        items.append((cat, title))

card_tpl = '''                <div class="product-card reveal">
                    <div class="product-img">
                        <img src="{img}" alt="{alt}" loading="lazy">
                    </div>
                    <div class="product-info">
                        <h3>{title}</h3>
                        <a href="contact.html" class="btn" style="padding: 10px 20px; font-size: 12px; margin-top: 15px;">Inquiry Now</a>
                    </div>
                </div>
'''

sections = []
total = 0
for key, name, desc, img in GROUPS:
    group = [t for (c, t) in items if c.lower() == key.lower()]
    if not group:
        continue
    total += len(group)
    cards = ''.join(
        card_tpl.format(img=img, alt=esc(t[:80]), title=esc(t))
        for t in group
    )
    sections.append('''    <section class="container" style="padding-top: 60px;">
        <div class="section-title reveal">
            <h2>{name}</h2>
            <p>{desc}</p>
            <p style="font-size: 13px; opacity: .6;">{n} products</p>
        </div>
        <div class="product-grid">
{cards}        </div>
    </section>
'''.format(name=esc(name), desc=esc(desc), n=len(group), cards=cards))

page = '''<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Product Catalog | Stark Sanitary - 100% Genuine Sanitary Ware</title>
    <meta name="description" content="Full product catalog of Nan'an Yingdu Stark Hardware Co., Ltd. - SUS304 stainless steel commercial sanitary ware, induction sensor sanitary ware and heavy-duty delay valve series.">
    <link rel="alternate" hreflang="en" href="../en/products.html" />
    <link rel="stylesheet" href="../assets/css/styles.css">
</head>
<body>

    <header>
        <div class="container">
            <nav>
                <a href="index.html" class="logo">STARK SANITARY</a>
                <ul class="nav-links">
                    <li><a href="index.html">Home</a></li>
                    <li><a href="about.html">About</a></li>
                    <li><a href="products.html">Products</a></li>
                    <li><a href="contact.html">Contact</a></li>
                </ul>
            </nav>
        </div>
    </header>

    <section class="container" style="padding-top: 150px;">
        <div class="section-title reveal">
            <h2>Product Catalog</h2>
            <p>{total} professional sanitary ware models - SUS304 Stainless Steel &amp; Solid Brass</p>
        </div>
    </section>

{sections}
    <footer>
        <div class="container">
            <div class="footer-grid">
                <div class="footer-col">
                    <h4>STARK SANITARY</h4>
                    <p>Professional Sanitary Ware Manufacturer since 2004.</p>
                </div>
                <div class="footer-col">
                    <h4>Quick Links</h4>
                    <ul>
                        <li><a href="index.html">Home</a></li>
                        <li><a href="products.html">Product Catalog</a></li>
                        <li><a href="contact.html">Request a Quote</a></li>
                    </ul>
                </div>
                <div class="footer-col">
                    <h4>Contact Us</h4>
                    <p>Email: dadadajiehui@gmail.com</p>
                    <p>Phone: 18522089987</p>
                    <p>Address: Area A, Phase 1, Luncang Plumbing City, Nan'an, Quanzhou, Fujian, China</p>
                </div>
            </div>
            <div class="copyright">
                <p>&copy; 2026 Nan'an Yingdu Stark Hardware Co., Ltd. All rights reserved.</p>
                <p style="font-size: 10px; margin-top: 10px; opacity: 0.5;">Powered by 郑在出海 | Warren | 微信：HK-1912</p>
            </div>
        </div>
    </footer>

    <script src="../assets/js/main.js"></script>
</body>
</html>
'''.format(total=total, sections='\n'.join(sections))

with open(OUT, 'w', encoding='utf-8', newline='\n') as f:
    f.write(page)

print('TOTAL_PRODUCTS=%d' % total)
print('BYTES=%d' % len(page.encode('utf-8')))
print('OUT=' + OUT)
