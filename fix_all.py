# -*- coding: utf-8 -*-
import os, re, html, zipfile

ROOT = r"C:\Users\董梅\AppData\Local\Temp\starkfix"
XLSX = r"C:\Users\董梅\AccioWork\2026-09-10-19-27-28-237-8204b69a\temp_data_package\资料准备说明\产品类目图片--中英文名字都对应下\产品图.xlsx"

SHOWER = "https://s.alicdn.com/@sc04/kf/H19c6adebd34f4d779b2e048e5952ad27G/SUS304-Stainless-Steel-Digital-Waterfall-Rain-Shower.jpg"
SENSOR = "https://s.alicdn.com/@sc04/kf/H9cad3bf47aaf46c79ad271e76469a458j/Chrome-Automatic-Sensor-Bathroom-Basin-Faucet-Touchless.jpg"
DELAY  = "https://s.alicdn.com/@sc04/kf/Hd87b785020ee4b56baa9b24c1d2957c4g/Double-Outlet-Garden-Tap-1-2-3.jpg"
KITCHEN = "https://s.alicdn.com/@sc04/kf/Hacea5b8e12b24dcb804ff5d197c689c5o/Matte-Black-High-Arc-Kitchen-Faucet-Pull.jpg"
THERMO = "https://s.alicdn.com/@sc04/kf/H1e6824632f7440a5acd7d1c2f296f314o/Modern-Matte-Black-Thermostatic-Shower-Mixer-Set.png"
TOILET = "https://s.alicdn.com/@sc04/kf/H99fbb47d32ed491d8a3267928eac2de3h/SUS304-Stainless-Steel-One-Piece-Floor-Mounted.jpg"

log = []

def rd(p):
    with open(p, "r", encoding="utf-8") as f:
        return f.read()

def wr(p, s):
    with open(p, "w", encoding="utf-8", newline="\n") as f:
        f.write(s)

def rep(s, old, new, label, required=True):
    if old not in s:
        if required:
            log.append("MISS: " + label)
            return s
        log.append("SKIP: " + label)
        return s
    n = s.count(old)
    log.append("OK(%d): %s" % (n, label))
    return s.replace(old, new)

# ---------------- zh/index.html ----------------
p = os.path.join(ROOT, "src", "zh", "index.html")
z = rd(p)
z = rep(z, '<a href="index.html" class="logo">斯塔克卫浴</a>',
           '<a href="index.html" class="logo">STARK SANITARY | 斯塔克卫浴</a>', "zh logo")
z = rep(z, '<section class="hero">',
           '<section class="hero" style="background-image: linear-gradient(rgba(10, 10, 10, 0.55), rgba(10, 10, 10, 0.68)), url(\'' + SHOWER + '\');">',
           "zh hero image")
z = rep(z, '<h1 class="reveal">灵动水生活，精工耀百年</h1>',
           '<h1 class="reveal">水之美学<br>精工智造</h1>', "zh h1")
z = rep(z, '<p class="reveal">20载匠心淬炼，深耕中国水暖之乡。专业智造高品质卫浴空间。</p>',
           '<p class="reveal">20载专注卫浴五金。深耕中国水暖之乡——南安仑苍。为您定义高端卫浴空间。</p>', "zh subtext")
# brand consistency in title
z = rep(z, '<title>斯塔克卫浴 | 专业卫浴水暖五金制造商</title>',
           '<title>STARK SANITARY 斯塔克卫浴 | 专业卫浴水暖五金制造商</title>', "zh title", required=False)
wr(p, z)

# ---------------- en/index.html ----------------
p = os.path.join(ROOT, "src", "en", "index.html")
e = rd(p)
e = rep(e, '<a href="index.html" class="logo">STARK HARDWARE</a>',
           '<a href="index.html" class="logo">STARK SANITARY</a>', "en logo")
e = rep(e, '<h4>STARK HARDWARE</h4>', '<h4>STARK SANITARY</h4>', "en footer brand")
e = rep(e, '<title>Stark Hardware | Professional Sanitary Ware & Plumbing Manufacturer</title>',
           '<title>Stark Sanitary | Professional Sanitary Ware & Plumbing Manufacturer</title>', "en title")
e = rep(e, '<h1 class="reveal">Elegance in Every Drop</h1>',
           '<h1 class="reveal">Elegance<br>In Every Drop</h1>', "en h1")

card1_old = ('<div class="product-img"><img src="' + KITCHEN + '" alt="Kitchen Faucet"></div>\n'
             '                    <div class="product-info"><h3>Kitchen Faucets</h3></div>')
card1_new = ('<div class="product-img"><img src="' + SHOWER + '" alt="Digital Shower Panel"></div>\n'
             '                    <div class="product-info"><h3>Digital Shower Panel</h3></div>')
e = rep(e, card1_old, card1_new, "en card1 Digital Shower Panel")

card2_old = ('<div class="product-img"><img src="' + THERMO + '" alt="Shower System"></div>\n'
             '                    <div class="product-info"><h3>Shower Systems</h3></div>')
card2_new = ('<div class="product-img"><img src="' + SENSOR + '" alt="Intelligent Sensor Faucet"></div>\n'
             '                    <div class="product-info"><h3>Intelligent Sensor Faucet</h3></div>')
e = rep(e, card2_old, card2_new, "en card2 Intelligent Sensor Faucet")

card3_old = ('<div class="product-img"><img src="' + SENSOR + '" alt="Basin Mixer"></div>\n'
             '                    <div class="product-info"><h3>Basin Mixers</h3></div>')
card3_new = ('<div class="product-img"><img src="' + DELAY + '" alt="Commercial Delay Valve"></div>\n'
             '                    <div class="product-info"><h3>Commercial Delay Valve</h3></div>')
e = rep(e, card3_old, card3_new, "en card3 Commercial Delay Valve")
wr(p, e)

# ---------------- assets/css/styles.css ----------------
p = os.path.join(ROOT, "src", "assets", "css", "styles.css")
c = rd(p)
c = rep(c, "linear-gradient(rgba(10, 10, 10, 0.55), rgba(10, 10, 10, 0.68)), var(--hero-image);",
           "linear-gradient(rgba(10, 10, 10, 0.55), rgba(10, 10, 10, 0.68)), url('" + SHOWER + "');",
           "css .hero direct url")
c = rep(c, ".logo {\n    font-size: 24px;",
           ".logo {\n    font-family: 'Inter', sans-serif;\n    font-size: 24px;",
           "css .logo font-family")
wr(p, c)

# ---------------- src/en/products.html ----------------
def read_rows(path):
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        out = []
        for r in ws.iter_rows(values_only=True):
            r = list(r) + ['', '', '', '']
            out.append({'A': r[0], 'B': r[1], 'C': r[2], 'D': r[3]})
        return out
    except Exception as ex:
        log.append("openpyxl-fallback: %s" % ex)
    rows = []
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        shared = []
        if 'xl/sharedStrings.xml' in names:
            s = zf.read('xl/sharedStrings.xml').decode('utf-8', 'replace')
            for si in re.findall(r'<si>(.*?)</si>', s, re.S):
                txt = ''.join(re.findall(r'<t[^>]*>(.*?)</t>', si, re.S))
                shared.append(html.unescape(txt))
        sheets = sorted(n for n in names if re.match(r'xl/worksheets/sheet\d+\.xml$', n))
        sheet = zf.read(sheets[0]).decode('utf-8', 'replace')
        for row in re.findall(r'<row[^>]*>(.*?)</row>', sheet, re.S):
            cells = {}
            for m in re.finditer(r'<c\b([^>]*?)(?:/>|>(.*?)</c>)', row, re.S):
                attrs = m.group(1) or ''
                body = m.group(2) or ''
                ref = re.search(r'r="([A-Z]+)\d+"', attrs)
                col = ref.group(1) if ref else 'A'
                t = re.search(r't="([^"]+)"', attrs)
                typ = t.group(1) if t else ''
                val = ''
                if typ == 'inlineStr':
                    val = ''.join(re.findall(r'<t[^>]*>(.*?)</t>', body, re.S))
                else:
                    v = re.search(r'<v>(.*?)</v>', body, re.S)
                    if v:
                        val = v.group(1)
                        if typ == 's':
                            try:
                                val = shared[int(val)]
                            except Exception:
                                pass
                cells[col] = html.unescape(val)
            rows.append(cells)
    return rows

def esc(s):
    s = str(s or '').replace('&amp;', '&').strip()
    return html.escape(s, quote=True)

GROUPS = [
    ("delay valve", "Commercial Delay Valve Series",
     "Heavy-duty brass self-closing and time-delay flush valves for public restrooms, hospitals and schools", DELAY),
    ("Induction Sanitary Ware Series", "Induction Sanitary Ware Series",
     "Touchless infrared sensor faucets, urinal flushers and automatic soap dispensers", SENSOR),
    ("304 stainless steel commercial series", "304 Stainless Steel Commercial Series",
     "SUS304 vandal-resistant sanitary ware for public, hotel, school and hospital projects", TOILET),
]

rows = read_rows(XLSX)
items = []
for r in rows[1:]:
    cat = str(r.get('B') or '').strip()
    title = str(r.get('C') or '').strip()
    if cat and title:
        items.append((cat, title))

card_tpl = '''                <div class="product-card reveal">
                    <div class="product-img">
                        <img src="{img}" alt="{alt}" loading="lazy">
                    </div>
                    <div class="product-info">
                        <h3>{title}</h3>
                        <a href="contact.html" class="btn" style="padding: 10px 20px; font-size: 12px; margin-top: 15px;">Inquiry Now</a>
                    </div>
                </div>
'''

sections = []
total = 0
for key, name, desc, img in GROUPS:
    group = [t for (cc, t) in items if cc.lower() == key.lower()]
    if not group:
        continue
    total += len(group)
    cards = ''.join(card_tpl.format(img=img, alt=esc(t[:80]), title=esc(t)) for t in group)
    sections.append('''    <section class="container" style="padding-top: 60px;">
        <div class="section-title reveal">
            <h2>{name}</h2>
            <p>{desc}</p>
            <p style="font-size: 13px; opacity: .6;">{n} products</p>
        </div>
        <div class="product-grid">
{cards}        </div>
    </section>
'''.format(name=esc(name), desc=esc(desc), n=len(group), cards=cards))

page = '''<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Product Catalog | Stark Sanitary - 100% Genuine Sanitary Ware</title>
    <meta name="description" content="Full product catalog of Nan'an Yingdu Stark Hardware Co., Ltd. - SUS304 stainless steel commercial sanitary ware, induction sensor sanitary ware and heavy-duty delay valve series.">
    <link rel="alternate" hreflang="en" href="../en/products.html" />
    <link rel="stylesheet" href="../assets/css/styles.css">
</head>
<body>

    <header>
        <div class="container">
            <nav>
                <a href="index.html" class="logo">STARK SANITARY</a>
                <ul class="nav-links">
                    <li><a href="index.html">Home</a></li>
                    <li><a href="about.html">About</a></li>
                    <li><a href="products.html">Products</a></li>
                    <li><a href="contact.html">Contact</a></li>
                </ul>
            </nav>
        </div>
    </header>

    <section class="container" style="padding-top: 150px;">
        <div class="section-title reveal">
            <h2>Product Catalog</h2>
            <p>{total} professional sanitary ware models - SUS304 Stainless Steel &amp; Solid Brass</p>
        </div>
    </section>

{sections}
    <footer>
        <div class="container">
            <div class="footer-grid">
                <div class="footer-col">
                    <h4>STARK SANITARY</h4>
                    <p>Professional Sanitary Ware Manufacturer since 2004.</p>
                </div>
                <div class="footer-col">
                    <h4>Quick Links</h4>
                    <ul>
                        <li><a href="index.html">Home</a></li>
                        <li><a href="products.html">Product Catalog</a></li>
                        <li><a href="contact.html">Request a Quote</a></li>
                    </ul>
                </div>
                <div class="footer-col">
                    <h4>Contact Us</h4>
                    <p>Email: dadadajiehui@gmail.com</p>
                    <p>Phone: 18522089987</p>
                    <p>Address: Area A, Phase 1, Luncang Plumbing City, Nan'an, Quanzhou, Fujian, China</p>
                </div>
            </div>
            <div class="copyright">
                <p>&copy; 2026 Nan'an Yingdu Stark Hardware Co., Ltd. All rights reserved.</p>
                <p style="font-size: 10px; margin-top: 10px; opacity: 0.5;">Powered by 郑在出海 | Warren | 微信：HK-1912</p>
            </div>
        </div>
    </footer>

    <script src="../assets/js/main.js"></script>
</body>
</html>
'''.format(total=total, sections='\n'.join(sections))

wr(os.path.join(ROOT, "src", "en", "products.html"), page)
log.append("PRODUCTS total=%d bytes=%d" % (total, len(page.encode('utf-8'))))

print("\n".join(log))
