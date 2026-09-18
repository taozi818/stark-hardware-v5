# -*- coding: utf-8 -*-
import os, re, html, zipfile, json, base64

ROOT = r"C:\Users\董梅\AppData\Local\Temp\starkfix"
OUT = r"C:\Users\董梅\AccioWork\2026-09-10-19-27-28-237-8204b69a"
XLSX = r"C:\Users\董梅\AccioWork\2026-09-10-19-27-28-237-8204b69a\temp_data_package\资料准备说明\产品类目图片--中英文名字都对应下\产品图.xlsx"
MSG = "CRITICAL: Content and Visual Upgrade v5.1"

SHOWER = "https://s.alicdn.com/@sc04/kf/H19c6adebd34f4d779b2e048e5952ad27G/SUS304-Stainless-Steel-Digital-Waterfall-Rain-Shower.jpg"
SENSOR = "https://s.alicdn.com/@sc04/kf/H9cad3bf47aaf46c79ad271e76469a458j/Chrome-Automatic-Sensor-Bathroom-Basin-Faucet-Touchless.jpg"
DELAY  = "https://s.alicdn.com/@sc04/kf/Hd87b785020ee4b56baa9b24c1d2957c4g/Double-Outlet-Garden-Tap-1-2-3.jpg"
TOILET = "https://s.alicdn.com/@sc04/kf/H99fbb47d32ed491d8a3267928eac2de3h/SUS304-Stainless-Steel-One-Piece-Floor-Mounted.jpg"

COMMIT = '''
  const changed = d !== before;
  if (changed) view.dispatch({ changes: { from: 0, to: view.state.doc.length, insert: d } });
  await new Promise(r => setTimeout(r, 1200));
  const pbtn = Array.from(document.querySelectorAll('button')).find(b => (b.textContent || '').trim() === '\u63d0\u4ea4\u66f4\u6539...');
  if (!pbtn) return JSON.stringify({ err: 'no page commit btn', changed: changed });
  pbtn.click();
  await new Promise(r => setTimeout(r, 3000));
  const input = document.getElementById('commit-message-input');
  if (!input) return JSON.stringify({ err: 'no msg input', changed: changed });
  const setter = Object.getOwnPropertyDescriptor(window.HTMLInputElement.prototype, 'value').set;
  setter.call(input, MSG);
  input.dispatchEvent(new Event('input', { bubbles: true }));
  const direct = document.querySelector('input[name="pr-choice"][value="direct"]');
  if (direct && !direct.checked) direct.click();
  await new Promise(r => setTimeout(r, 800));
  const cbtn = Array.from(document.querySelectorAll('button')).find(b => (b.textContent || '').trim() === '\u63d0\u4ea4\u66f4\u6539');
  if (!cbtn) return JSON.stringify({ err: 'no modal commit btn', changed: changed });
  cbtn.click();
  await new Promise(r => setTimeout(r, 9000));
  return JSON.stringify({ changed: changed, url: location.href });
})()
'''


def b64_script(path, outname):
    with open(path, 'rb') as f:
        b = base64.b64encode(f.read()).decode('ascii')
    s = ("(async () => {\n"
         "  const MSG = " + json.dumps(MSG) + ";\n"
         "  const B = " + json.dumps(b) + ";\n"
         "  const bin = atob(B);\n"
         "  const arr = new Uint8Array(bin.length);\n"
         "  for (let i = 0; i < bin.length; i++) arr[i] = bin.charCodeAt(i);\n"
         "  const d = new TextDecoder('utf-8').decode(arr);\n"
         "  const view = document.querySelector('.cm-content').cmTile.view;\n"
         "  const before = view.state.doc.toString();\n" + COMMIT)
    with open(os.path.join(OUT, outname), 'w', encoding='utf-8', newline='\n') as f:
        f.write(s)
    return len(s)


def read_rows(path):
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        out = []
        for r in ws.iter_rows(values_only=True):
            r = list(r) + ['', '', '', '']
            out.append({'B': r[1], 'C': r[2]})
        return out
    except Exception:
        pass
    rows = []
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        shared = []
        if 'xl/sharedStrings.xml' in names:
            s = zf.read('xl/sharedStrings.xml').decode('utf-8', 'replace')
            for si in re.findall(r'<si>(.*?)</si>', s, re.S):
                txt = ''.join(re.findall(r'<t[^>]*>(.*?)</t>', si, re.S))
                shared.append(html.unescape(txt))
        sheets = sorted(n for n in names if re.match(r'xl/worksheets/sheet\d+\.xml$', n))
        sheet = zf.read(sheets[0]).decode('utf-8', 'replace')
        for row in re.findall(r'<row[^>]*>(.*?)</row>', sheet, re.S):
            cells = {}
            for m in re.finditer(r'<c\b([^>]*?)(?:/>|>(.*?)</c>)', row, re.S):
                attrs = m.group(1) or ''
                body = m.group(2) or ''
                ref = re.search(r'r="([A-Z]+)\d+"', attrs)
                col = ref.group(1) if ref else 'A'
                t = re.search(r't="([^"]+)"', attrs)
                typ = t.group(1) if t else ''
                val = ''
                if typ == 'inlineStr':
                    val = ''.join(re.findall(r'<t[^>]*>(.*?)</t>', body, re.S))
                else:
                    v = re.search(r'<v>(.*?)</v>', body, re.S)
                    if v:
                        val = v.group(1)
                        if typ == 's':
                            try:
                                val = shared[int(val)]
                            except Exception:
                                pass
                cells[col] = html.unescape(val)
            rows.append(cells)
    return rows


rows = read_rows(XLSX)
items = []
for r in rows[1:]:
    cat = str(r.get('B') or '').strip()
    title = str(r.get('C') or '').strip()
    if cat and title:
        items.append([cat, title])

GROUPS = [
    ["delay valve", "Commercial Delay Valve Series",
     "Heavy-duty brass self-closing and time-delay flush valves for public restrooms, hospitals and schools", DELAY],
    ["Induction Sanitary Ware Series", "Induction Sanitary Ware Series",
     "Touchless infrared sensor faucets, urinal flushers and automatic soap dispensers", SENSOR],
    ["304 stainless steel commercial series", "304 Stainless Steel Commercial Series",
     "SUS304 vandal-resistant sanitary ware for public, hotel, school and hospital projects", TOILET],
]

gdata = []
total = 0
for key, name, desc, img in GROUPS:
    g = [t for (c, t) in items if c.lower() == key.lower()]
    total += len(g)
    gdata.append({"n": name, "d": desc, "img": img, "t": g})

HEADER = '''<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Product Catalog | Stark Sanitary - 100% Genuine Sanitary Ware</title>
    <meta name="description" content="Full product catalog of Nan'an Yingdu Stark Hardware Co., Ltd. - SUS304 stainless steel commercial sanitary ware, induction sensor sanitary ware and heavy-duty delay valve series.">
    <link rel="alternate" hreflang="en" href="../en/products.html" />
    <link rel="stylesheet" href="../assets/css/styles.css">
</head>
<body>

    <header>
        <div class="container">
            <nav>
                <a href="index.html" class="logo">STARK SANITARY</a>
                <ul class="nav-links">
                    <li><a href="index.html">Home</a></li>
                    <li><a href="about.html">About</a></li>
                    <li><a href="products.html">Products</a></li>
                    <li><a href="contact.html">Contact</a></li>
                </ul>
            </nav>
        </div>
    </header>

    <section class="container" style="padding-top: 150px;">
        <div class="section-title reveal">
            <h2>Product Catalog</h2>
            <p>__TOTAL__ professional sanitary ware models - SUS304 Stainless Steel &amp; Solid Brass</p>
        </div>
    </section>

'''

FOOTER = '''    <footer>
        <div class="container">
            <div class="footer-grid">
                <div class="footer-col">
                    <h4>STARK SANITARY</h4>
                    <p>Professional Sanitary Ware Manufacturer since 2004.</p>
                </div>
                <div class="footer-col">
                    <h4>Quick Links</h4>
                    <ul>
                        <li><a href="index.html">Home</a></li>
                        <li><a href="products.html">Product Catalog</a></li>
                        <li><a href="contact.html">Request a Quote</a></li>
                    </ul>
                </div>
                <div class="footer-col">
                    <h4>Contact Us</h4>
                    <p>Email: dadadajiehui@gmail.com</p>
                    <p>Phone: 18522089987</p>
                    <p>Address: Area A, Phase 1, Luncang Plumbing City, Nan'an, Quanzhou, Fujian, China</p>
                </div>
            </div>
            <div class="copyright">
                <p>&copy; 2026 Nan'an Yingdu Stark Hardware Co., Ltd. All rights reserved.</p>
                <p style="font-size: 10px; margin-top: 10px; opacity: 0.5;">Powered by \u90d1\u5728\u51fa\u6d77 | Warren | \u5fae\u4fe1\uff1aHK-1912</p>
            </div>
        </div>
    </footer>

    <script src="../assets/js/main.js"></script>
</body>
</html>
'''

# products builder script
ps = ("(async () => {\n"
      "  const MSG = " + json.dumps(MSG) + ";\n"
      "  const G = " + json.dumps(gdata, ensure_ascii=False) + ";\n"
      "  const H = " + json.dumps(HEADER.replace('__TOTAL__', str(total)), ensure_ascii=False) + ";\n"
      "  const F = " + json.dumps(FOOTER, ensure_ascii=False) + ";\n"
      "  const esc = s => String(s).replace(/&/g,'&amp;').replace(/</g,'&lt;').replace(/>/g,'&gt;').replace(/\"/g,'&quot;');\n"
      "  const card = (img, arr) => arr.map(t => '                <div class=\"product-card reveal\">\\n'\n"
      "      + '                    <div class=\"product-img\">\\n'\n"
      "      + '                        <img src=\"' + img + '\" alt=\"' + esc(t.slice(0,80)) + '\" loading=\"lazy\">\\n'\n"
      "      + '                    </div>\\n'\n"
      "      + '                    <div class=\"product-info\">\\n'\n"
      "      + '                        <h3>' + esc(t) + '</h3>\\n'\n"
      "      + '                        <a href=\"contact.html\" class=\"btn\" style=\"padding: 10px 20px; font-size: 12px; margin-top: 15px;\">Inquiry Now</a>\\n'\n"
      "      + '                    </div>\\n'\n"
      "      + '                </div>\\n').join('');\n"
      "  const secs = G.map(g => '    <section class=\"container\" style=\"padding-top: 60px;\">\\n'\n"
      "      + '        <div class=\"section-title reveal\">\\n'\n"
      "      + '            <h2>' + esc(g.n) + '</h2>\\n'\n"
      "      + '            <p>' + esc(g.d) + '</p>\\n'\n"
      "      + '            <p style=\"font-size: 13px; opacity: .6;\">' + g.t.length + ' products</p>\\n'\n"
      "      + '        </div>\\n'\n"
      "      + '        <div class=\"product-grid\">\\n'\n"
      "      + card(g.img, g.t)\n"
      "      + '        </div>\\n'\n"
      "      + '    </section>\\n').join('');\n"
      "  const d = H + secs + F;\n"
      "  let view = (document.querySelector('.cm-content') || {}).cmTile;\n"
      "  view = view ? view.view : null;\n"
      "  if (!view) return JSON.stringify({ err: 'no codemirror', len: d.length });\n"
      "  const before = view.state.doc.toString();\n"
      "  view.dispatch({ changes: { from: 0, to: before.length, insert: '' } });\n"
      "  view.dispatch({ changes: { from: 0, to: 0, insert: d } });\n"
      + COMMIT)

with open(os.path.join(OUT, "gh_products.js"), "w", encoding="utf-8", newline="\n") as f:
    f.write(ps)

print("products total=%d js_bytes=%d" % (total, len(ps.encode('utf-8'))))
print("zh=%d en=%d css=%d" % (
    b64_script(os.path.join(ROOT, "src", "zh", "index.html"), "gh_zh.js"),
    b64_script(os.path.join(ROOT, "src", "en", "index.html"), "gh_en.js"),
    b64_script(os.path.join(ROOT, "src", "assets", "css", "styles.css"), "gh_css.js")))
